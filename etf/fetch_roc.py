"""抓取发行商官网的"返还本金（Return of Capital, ROC）"数据，输出 etf/raw/roc_*.csv。

来源（全部是发行商官网原始文件）：
  SPYI / QQQI：NEOS 官网 Form 8937（年度最终税务数据，逐次分红的 ROC 金额）+ 最新 19a-1 通知（当年估计）
  QYLD       ：Global X 官网 Form 8937（财年截至 10/31）+ 年度 1099 补充表（日历年）+ 最新 19a 通知（估计）
  JEPI / JEPQ：J.P. Morgan 官网分红表、年度分红通知里都没有 ROC 字段 -> 未取得
部分老 8937 是扫描图片，已人工对照图片转录到 raw/roc_transcribed_from_scans.csv。

用法：python etf/fetch_roc.py   依赖：requests pypdf python-docx pandas
"""
import io
import re
from pathlib import Path

import pandas as pd
import pypdf
import requests

RAW = Path(__file__).parent / "raw"
PDF = RAW / "issuer_docs"
PDF.mkdir(parents=True, exist_ok=True)
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124.0 Safari/537.36"}
NEOS = "https://neosfunds.com/wp-content/uploads/"
GX = "https://assets.globalxetfs.com/funds/tax_supplements/"

NEOS_8937 = {  # 可提取文字的 8937（扫描件见转录 CSV）
    ("SPYI", "FYE 2025-05-31"): NEOS + "SPYI-Form-8937-5.31.25.pdf",
    ("SPYI", "FYE 2025-12-31"): NEOS + "NEOS-SP-500-R-High-Income-ETF-Form-8937-12.31.25.pdf",
    ("QQQI", "FYE 2025-05-31"): NEOS + "QQQI-Form-8937-5.31.25.pdf",
    ("QQQI", "FYE 2025-12-31"): NEOS + "NEOS-Nasdaq-100-High-Income-ETF-Form-8937-12.31.25.pdf",
}
NEOS_19A_LATEST = {"SPYI": NEOS + "SPYI-19a-1-Notice-9.16.26-Confidential.pdf",
                   "QQQI": NEOS + "QQQI-19a-1-Notice-9.16.26-Confidential.pdf"}
QYLD_8937 = {y: GX + f"QYLD_Form-8937_1031{y}.pdf" for y in (2019, 2020, 2022, 2023, 2025)}
GX_YEAR_END = {y: f"https://assets-cms.globalxetfs.com/{y}-Year-End-Tax-Supplement-Global-X-ETFs.pdf"
               for y in (2023, 2024, 2025)}
QYLD_19A_LATEST = GX + "QYLD_Form-19a_09242026.docx"


def get(url):
    name = PDF / url.rsplit("/", 1)[-1]
    if not name.exists():
        r = requests.get(url, headers=UA, timeout=90)
        r.raise_for_status()
        name.write_bytes(r.content)
    return name


def pdf_text(url):
    return "\n".join(p.extract_text() or "" for p in pypdf.PdfReader(get(url)).pages)


def iso(d):
    m, dd, y = d.split("/")
    return f"{y}-{int(m):02d}-{int(dd):02d}"


def neos_8937():
    rows = []
    row = re.compile(r"(\d+/\d+/\d{4})\s+(\d+/\d+/\d{4})\s+(\d+/\d+/\d{4})\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)%")
    for (fund, period), url in NEOS_8937.items():
        for m in row.finditer(pdf_text(url)):
            rows.append({"Fund": fund, "Source": "Form 8937", "Period": period, "ExDate": iso(m[2]),
                         "Amount": float(m[4]), "ROC_Amount": float(m[5]), "ROC_Pct": float(m[6]),
                         "Note": url})
    return rows


def qyld_8937():
    """2019 版列序：金额, ROC%, ROC额, 应税%, …；2020 起：金额, 应税%, 应税额, ROC%, ROC额（已对照原件图片确认）。"""
    rows = []
    for y, url in QYLD_8937.items():
        for ln in pdf_text(url).splitlines():
            m = re.match(r"(\d+/\d+/\d{4})\s+\d+/\d+/\d{4}\s+\d+/\d+/\d{4}\s+([\d.]+)\s+(.*)", ln.strip())
            if not m:
                continue
            nums = re.findall(r"([\d.]+)%\s+([\d.]+)", m[3])
            pct, amt = (nums[0] if y == 2019 else nums[1])
            rows.append({"Fund": "QYLD", "Source": "Form 8937", "Period": f"FYE {y}-10-31",
                         "ExDate": iso(m[1]), "Amount": float(m[2]), "ROC_Amount": float(amt),
                         "ROC_Pct": float(pct), "Note": url})
    return rows


def qyld_year_end():
    """Global X 日历年 1099 补充表：第 1 列=总分红，第 11 列=Return of Capital（ROC）。"""
    rows, seen = [], set()
    for y, url in GX_YEAR_END.items():
        txt = pdf_text(url)
        # 每个文件可能含多个年份的段落，按 "Calendar Year XXXX" 切段
        parts = re.split(r"for\s+Calendar\s+Year\s+(\d{4})", txt)
        for cy, body in zip(parts[1::2], parts[2::2]):
            flat = re.sub(r"\s+", "", body)  # 去掉所有空白（含换行、不间断空格），数字靠 $ 分隔
            for m in re.finditer(r"QYLD(\d\d/\d\d/\d{4})(\d\d/\d\d/\d{4})(\d\d/\d\d/\d{4})((?:[\d.]+\$){11})", flat):
                vals = [float(v) for v in re.findall(r"([\d.]+)\$", m[4])]
                key = (cy, m[2])
                if key in seen:
                    continue
                seen.add(key)
                rows.append({"Fund": "QYLD", "Source": "年度1099补充表", "Period": f"日历年 {cy}",
                             "ExDate": iso(m[2]), "Amount": vals[0], "ROC_Amount": vals[10],
                             "ROC_Pct": round(vals[10] / vals[0] * 100, 4) if vals[0] else 0, "Note": url})
    return rows


def latest_19a():
    out = []
    for fund, url in NEOS_19A_LATEST.items():
        t = re.sub(r"\s+", " ", pdf_text(url))
        m = re.search(r"Estimated Return of Capital \$([\d.]+) (\d+)% \$([\d.]+) (\d+)%", t)
        tot = re.search(r"Total \(per common share\) \$([\d.]+) 100% \$([\d.]+) 100%", t)
        out.append({"Fund": fund, "Source": "19a-1 通知（估计，非最终税务数据）", "Period": "2026 财年至今（截至 2026-09 分红）",
                    "Amount_YTD": float(tot[2]), "ROC_Amount_YTD": float(m[3]), "ROC_Pct_YTD": float(m[4]),
                    "Note": url})
    try:
        import docx
        d = docx.Document(get(QYLD_19A_LATEST))
        cells = [[c.text.strip() for c in r.cells] for t in d.tables for r in t.rows]
        text = " ".join(p.text for p in d.paragraphs)
        roc = [r for r in cells if any("Return of Capital" in c for c in r)]
        tot = [r for r in cells if r[0].lower().startswith("total (per")]
        nums = lambda r: [float(x) for c in r for x in re.findall(r"[\d.]+", c.replace(",", ""))]
        rn, tn = nums(roc[0]), nums(tot[0])
        # 表格列：本次金额, 本次%, 财年累计金额, 财年累计%
        out.append({"Fund": "QYLD", "Source": "19a 通知（估计，非最终税务数据）", "Period": "2026 财年至今（截至 2026-09 分红）",
                    "Amount_YTD": tn[2], "ROC_Amount_YTD": rn[2], "ROC_Pct_YTD": rn[3], "Note": QYLD_19A_LATEST})
    except Exception as e:  # noqa: BLE001
        out.append({"Fund": "QYLD", "Source": "19a 通知", "Period": "2026 财年至今", "Note": f"解析失败：{e}"})
    return out


def main():
    rows = neos_8937() + qyld_8937() + qyld_year_end()
    scans = pd.read_csv(RAW / "roc_transcribed_from_scans.csv")
    df = pd.concat([pd.DataFrame(rows), scans], ignore_index=True).sort_values(["Fund", "ExDate"])
    df.to_csv(RAW / "roc_per_distribution.csv", index=False)
    pd.DataFrame(latest_19a()).to_csv(RAW / "roc_19a_latest.csv", index=False)
    print(df.groupby(["Fund", "Source", "Period"]).agg(n=("Amount", "size"), amt=("Amount", "sum"),
                                                        roc=("ROC_Amount", "sum")).assign(pct=lambda x: x.roc / x.amt))
    print(pd.read_csv(RAW / "roc_19a_latest.csv").to_string())


if __name__ == "__main__":
    main()
